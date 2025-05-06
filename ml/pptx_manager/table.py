from loguru import logger
from pptx import Presentation
from pptx.util import Inches


class TableManager:
    def __init__(self) -> None:
        self.pres = None

    def load_presentation(self, file_path: str) -> str:
        self.pres = Presentation(file_path)
        logger.info(f'Загружена презентация: {file_path}')
        return f'Презентация успешно загружена: {file_path}'

    def get_table_indexes(self, slide_num: int) -> list[int]:
        if not self.pres:
            msg = 'Презентация не загружена. Сначала используйте `load_presentation`.'
            raise ValueError(msg)

        if not (1 <= slide_num <= len(self.pres.slides)):
            msg = f'Неверный номер слайда: {slide_num}'
            raise ValueError(msg)

        slide = self.pres.slides[slide_num - 1]
        return [i for i, shape in enumerate(slide.shapes) if shape.has_table]

    def parse_tables(self) -> list[dict]:
        if not self.pres:
            msg = 'Презентация не загружена. Сначала используйте `load_presentation`.'
            raise ValueError(msg)

        tables = []
        for slide_num, slide in enumerate(self.pres.slides, 1):
            table_indexes = self.get_table_indexes(slide_num)
            for table_idx in table_indexes:
                table = slide.shapes[table_idx].table
                tables.append(
                    {
                        'slide_num': slide_num,
                        'table_index': table_idx,
                        'data': [[cell.text.strip() for cell in row.cells] for row in table.rows],
                    }
                )
                logger.debug(f'Parsed table on slide {slide_num}, index {table_idx}')

        logger.info(f'Total tables parsed: {len(tables)}')
        return tables

    def create_table(
        self,
        slide_num: int,
        rows: int,
        cols: int,
        left: float = 1.0,
        top: float = 1.0,
        width: float = 6.0,
        height: float = 4.0,
    ) -> str:
        """Создает новую таблицу на указанном слайде.

        Создает таблицу с заданными размерами и позиционированием на слайде.

        Args:
            slide_num (int): Номер слайда (нумерация с 1).
            rows (int): Количество строк в таблице.
            cols (int): Количество столбцов в таблице.
            left (float, optional): Отступ слева в дюймах. По умолчанию 1.0.
            top (float, optional): Отступ сверху в дюймах. По умолчанию 1.0.
            width (float, optional): Ширина таблицы в дюймах. По умолчанию 6.0.
            height (float, optional): Высота таблицы в дюймах. По умолчанию 4.0.

        Returns:
            str: Сообщение о выполненных изменениях, включая ID слайда, фигур, и измененных параметрах.

        Raises:
            ValueError: Если презентация не загружена или параметры неверны.
        """
        logger.debug(f'Создание таблицы: слайд={slide_num}, строки={rows}, столбцы={cols}')

        if not self.pres:
            msg = 'Презентация не загружена. Используйте метод load_presentation().'
            logger.error(msg)
            raise ValueError(msg)

        if not (1 <= slide_num <= len(self.pres.slides)):
            msg = f'Недопустимый номер слайда: {slide_num}. Доступны слайды: 1-{len(self.pres.slides)}'
            logger.error(msg)
            raise ValueError(msg)

        slide = self.pres.slides[slide_num - 1]
        _ = slide.shapes.add_table(rows, cols, Inches(left), Inches(top), Inches(width), Inches(height))
        table_index = len(slide.shapes) - 1

        logger.info(f'Создана таблица {rows}x{cols} на слайде {slide_num}')
        return f'Создана таблица размером {rows}x{cols} на слайде {slide_num} (индекс: {table_index})'

    def update_cell(self, slide_num: int, table_index: int, row: int, col: int, text: str) -> str:
        """Обновляет содержимое определенной ячейки в таблице.

        Изменяет текстовое содержимое указанной ячейки в таблице на заданном слайде.

        Args:
            slide_num (int): Номер слайда (нумерация с 1).
            table_index (int): Индекс таблицы на слайде.
            row (int): Номер строки (нумерация с 0).
            col (int): Номер столбца (нумерация с 0).
            text (str): Новое содержимое для ячейки.

        Returns:
            str: Success message.

        Raises:
            ValueError: Если презентация не загружена или параметры недействительны.
            IndexError: Если индексы строки или столбца выходят за пределы таблицы.
        """
        logger.debug(f'Обновление ячейки: слайд={slide_num}, таблица={table_index}, строка={row}, столбец={col}')

        if not self.pres:
            msg = 'Презентация не загружена. Используйте метод load_presentation()'
            logger.error(msg)
            raise ValueError(msg) from None

        slide = self.pres.slides[slide_num - 1]
        shape = slide.shapes[table_index]

        if not shape.has_table:
            msg = f'Фигура с индексом {table_index} не является таблицей'
            logger.error(msg)
            raise ValueError(msg)

        table = shape.table

        if row >= len(table.rows) or col >= len(table.columns):
            msg = f'Индексы выходят за пределы таблицы: строка={row}, столбец={col}'
            logger.error(msg)
            raise IndexError(msg)

        cell = table.cell(row, col)
        old_text = cell.text_frame.text
        cell.text_frame.clear()
        cell.text_frame.paragraphs[0].text = text

        logger.info(f'Обновлена ячейка [{row},{col}] в таблице {table_index} на слайде {slide_num}')
        return (
            f'На слайде {slide_num} в таблице {table_index} обновлена ячейка [{row},{col}]. '
            f'Старый текст: "{old_text}", новый текст: "{text}"'
        )

    def save_presentation(self, file_path: str) -> str:
        if not self.pres:
            msg = 'Презентация не загружена. Используйте метод load_presentation()'
            logger.error(msg)
            raise ValueError(msg)

        self.pres.save(file_path)
        logger.info(f'Презентация сохранена: {file_path}')
        return f'Презентация успешно сохранена в файл: {file_path}'

    def test(self, source: str | None) -> None:
        try:
            self.load_presentation(source)

            existing_tables = self.parse_tables()

            logger.info(f'Found {len(existing_tables)} existing tables')

            new_table_message = self.create_table(1, 3, 4)
            logger.info(new_table_message)

            if existing_tables:
                first_table = existing_tables[0]
                update_message = self.update_cell(
                    first_table['slide_num'],
                    first_table['table_index'],
                    0,
                    0,
                    'Обновились или я тебя удалю, это будет больно',
                )
                logger.info(update_message)

            update_message = self.update_cell(1, len(self.pres.slides[0].shapes) - 1, 0, 0, 'New Table Content')
            logger.info(update_message)

            save_message = self.save_presentation('test_table.pptx')
            logger.info(save_message)

            logger.success('Тест успешно завершен.')

        except Exception as e:
            logger.error(f'Тест завершился с ошибкой: {str(e)}')


if __name__ == '__main__':
    tm = TableManager()
    tm.test('../test_data/test_dit.pptx')
